"""
Build the result matrix workbook from the main rates file.

Layout (matches rate-card matrix style):
  - Rows 1-3: headers
      * shipment columns: title on row 1 (rows 2-3 blank)
      * each cost = 2- or 3-column block:
          row1: Cost name          (merged)
          row2: Rate by: {basis}   (merged)
          row3: Currency | [MIN Flat if any] | p/unit or Flat
            - Rate by "per shipment" -> amount column header "Flat"
            - MIN Flat column omitted when no minimums exist
  - Row 4+: lane data
      currency | [minimum] | rate amount

Cost sources in main rates:
  - Base Rate (+ Currency / Basis / Minimum)
  - Every "* Value" / "* Details" pair
      Details "USD, per w/m" -> currency=USD, rate by=per w/m
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import paths as project_paths
from city_naming import (
    apply_city_rename,
    build_lane_city_index,
    city_match_key,
    format_city_with_aliases,
    lookup_city_spelling,
)

ROOT = project_paths.ROOT
PROCESSING_DIR = project_paths.PROCESSING_DIR
OUTPUT_DIR = project_paths.OUTPUT_DIR


def _sync_dirs() -> None:
    global ROOT, PROCESSING_DIR, OUTPUT_DIR
    ROOT = project_paths.ROOT
    PROCESSING_DIR = project_paths.PROCESSING_DIR
    OUTPUT_DIR = project_paths.OUTPUT_DIR

DATA_START_ROW = 4

SHIPMENT_HEADERS = [
    "Lane #",
    "Rate Card",
    "UNIQUE_IDENTIFIER",
    "Commodity",
    "Origin City",
    "Origin City",
    "Origin Country",
    "Destination City",
    "Destination City",
    "Destination Country",
    "Valid from",
    "Valid to",
    "Transport Mode",
]

# Source cost name -> display name in the result matrix
COST_DISPLAY_NAMES = {
    "Bunker Adjustment Factor": "BAF Fee",
    "Destination Handling": "Destination Terminal Handling Fee",
    "EU ETS": "EU ETS Fee",
    "Hazardous": "Hazardous Fee",
    "Origin Handling": "Origin Terminal Handling Fee",
}


def _strip_columns(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(c).strip() for c in out.columns]
    return out


def _before_comma(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    if not text:
        return None
    return text.split(",", 1)[0].strip()


def _commodity_after_dash(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    if " - " in text:
        return text.split(" - ", 1)[1].strip()
    if "-" in text:
        return text.split("-", 1)[1].strip()
    return text


def _format_date(value) -> str | None:
    if pd.isna(value):
        return None
    ts = pd.to_datetime(value, errors="coerce")
    if pd.isna(ts):
        return str(value)
    return ts.strftime("%d.%m.%Y")


def format_unique_identifier(value) -> str | None:
    """Keep Unique Identifier as the full number shown in the cell (no sci notation)."""
    if pd.isna(value):
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        # Already a plain digit string
        if re.fullmatch(r"\d+", text):
            return text
        try:
            value = float(text)
        except ValueError:
            return text
    number = float(value)
    if number.is_integer():
        return str(int(number))
    return str(visible_number(number))


def visible_number(value):
    """Match Excel-visible numeric display (drop binary float noise)."""
    if pd.isna(value):
        return None
    if isinstance(value, str):
        cleaned = value.strip()
        if not cleaned:
            return None
        match = re.search(r"-?\d+(?:\.\d+)?", cleaned.replace(",", ""))
        if not match:
            return None
        value = float(match.group(0))

    number = float(value)
    rounded = round(number + (1e-12 if number >= 0 else -1e-12), 2)
    if rounded == int(rounded):
        return float(int(rounded))
    return rounded


def parse_minimum_amount(value):
    if pd.isna(value):
        return None
    text = str(value).strip()
    if not text:
        return None
    match = re.search(r"-?\d+(?:,\d{3})*(?:\.\d+)?", text)
    if not match:
        return None
    return visible_number(match.group(0).replace(",", ""))


def rate_unit_label(basis) -> str:
    """Map rate basis to p/unit or Flat (rate amount column header)."""
    if pd.isna(basis) or not str(basis).strip():
        return "p/unit"
    text = str(basis).strip().lower()
    # per shipment (and similar) -> Flat instead of p/unit
    flat_tokens = ("flat", "shipment", "lump", "fixed", "per container", "per ctr")
    if any(token in text for token in flat_tokens):
        return "Flat"
    return "p/unit"


def cost_has_minimum(minimums: list) -> bool:
    """True when at least one lane has a MIN value."""
    return any(m is not None for m in minimums)


def parse_details(details) -> tuple[str | None, str | None]:
    """Parse Details like 'USD, per w/m' -> (currency, rate_by)."""
    if pd.isna(details):
        return None, None
    text = str(details).strip()
    if not text:
        return None, None
    if "," in text:
        currency, basis = text.split(",", 1)
        return currency.strip() or None, basis.strip() or None
    return text, None


def main_file_stem(path: Path) -> str:
    """Stem of the main rates file, without _processed suffix."""
    stem = path.stem
    if stem.endswith("_processed"):
        stem = stem[: -len("_processed")]
    return stem


def rate_card_name_from_path(path: Path) -> str:
    return f"{main_file_stem(path)}.xlsx"


def result_output_path(main_rates_path: Path, rate_card: str | None = None) -> Path:
    """output/{main_file_name}_result.xlsx"""
    _sync_dirs()
    if rate_card:
        stem = Path(rate_card).stem
    else:
        stem = main_file_stem(Path(main_rates_path))
    return OUTPUT_DIR / f"{stem}_result.xlsx"


def load_main_rates(path: Path) -> pd.DataFrame:
    frames = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    if "Mainleg rates" in frames:
        df = frames["Mainleg rates"]
    else:
        df = next(iter(frames.values()))
    return _strip_columns(df)


def build_shipment_rows(df: pd.DataFrame, rate_card: str) -> list[list]:
    """
    Build Result/BAF/Missing shipment rows.
    Origin/Destination cities: rename table + common-rating aliases
    (same duplicate-safe rules as Accessorial).
    """
    rows_meta: list[dict] = []
    has_hazardous = "Hazardous Value" in df.columns

    for _, row in df.iterrows():
        origin_raw = (
            "" if pd.isna(row.get("Origin City")) else str(row.get("Origin City")).strip()
        )
        dest_raw = (
            ""
            if pd.isna(row.get("Destination City"))
            else str(row.get("Destination City")).strip()
        )

        origin_base, origin_aliases = _resolve_accessorial_city(
            origin_raw, strip_comma=True
        )
        dest_base, dest_aliases = _resolve_accessorial_city(dest_raw, strip_comma=True)

        # Preserve destination region/state suffix for the "full" city column
        dest_suffix = ""
        if dest_raw and "," in dest_raw:
            _left, right = dest_raw.split(",", 1)
            # Keep suffix only when rename did not consume the comma form
            renamed_full = apply_city_rename(dest_raw)[0]
            if renamed_full.casefold() == dest_raw.casefold():
                dest_suffix = right.strip()

        commodity = _commodity_after_dash(row["Commodity"])
        if has_hazardous and visible_number(row.get("Hazardous Value")) is not None:
            commodity = f"{commodity}, HAZ" if commodity else "HAZ"

        rows_meta.append(
            {
                "origin": origin_base,
                "origin_aliases": origin_aliases,
                "dest": dest_base,
                "dest_aliases": dest_aliases,
                "dest_suffix": dest_suffix,
                "origin_service": _strip_service(row.get("Origin Service")),
                "dest_service": _strip_service(row.get("Destination Service")),
                "uid": format_unique_identifier(row["Unique Identifier"]),
                "commodity": commodity,
                "origin_country": (
                    None
                    if pd.isna(row["Origin Country"])
                    else str(row["Origin Country"]).strip()
                ),
                "dest_country": (
                    None
                    if pd.isna(row["Destination Country"])
                    else str(row["Destination Country"]).strip()
                ),
                "valid_from": _format_date(row["Line Item Eff Date"]),
                "valid_to": _format_date(row["Line Item Exp Date"]),
            }
        )

    lane_index = build_lane_city_index(
        [
            (m["origin"], m["origin_service"], m["dest"], m["dest_service"])
            for m in rows_meta
        ]
    )

    rows: list[list] = []
    for m in rows_meta:
        from_key = city_match_key(m["origin"])
        to_key = city_match_key(m["dest"])
        origin_city = format_city_with_aliases(
            m["origin"],
            side="from",
            from_key=from_key,
            from_service=m["origin_service"],
            to_key=to_key,
            to_service=m["dest_service"],
            lane_index=lane_index,
            rename_aliases=m["origin_aliases"],
        )
        dest_city = format_city_with_aliases(
            m["dest"],
            side="to",
            from_key=from_key,
            from_service=m["origin_service"],
            to_key=to_key,
            to_service=m["dest_service"],
            lane_index=lane_index,
            rename_aliases=m["dest_aliases"],
        )
        dest_full = (
            f"{dest_city}, {m['dest_suffix']}" if m["dest_suffix"] else dest_city
        )
        rows.append(
            [
                len(rows) + 1,
                rate_card,
                m["uid"],
                m["commodity"],
                origin_city,
                origin_city,
                m["origin_country"],
                dest_full,
                dest_city,
                m["dest_country"],
                m["valid_from"],
                m["valid_to"],
                "OCEAN",
            ]
        )
    return rows


def _shared_or_first(values: list[str | None]) -> str:
    cleaned = [str(v).strip() for v in values if v is not None and str(v).strip()]
    if not cleaned:
        return ""
    unique = list(dict.fromkeys(cleaned))
    return unique[0]


def build_base_rate_cost(df: pd.DataFrame) -> dict:
    """Base Rate as a 3-col block: Currency | MIN Flat | p/unit|flat."""
    currencies: list = []
    minimums: list = []
    rates: list = []
    bases: list[str] = []

    for _, row in df.iterrows():
        basis = row.get("Base Rate Basis")
        basis_text = "" if pd.isna(basis) else str(basis).strip()
        currency = row.get("Base Rate Currency")
        currency_text = (
            None if pd.isna(currency) else str(currency).strip() or None
        )

        currencies.append(currency_text)
        minimums.append(parse_minimum_amount(row.get("Minimum")))
        rates.append(visible_number(row.get("Base Rate")))
        bases.append(basis_text)

    rate_by = _shared_or_first(bases)
    return {
        "cost_name": "Base Rate",
        "rate_by": rate_by,
        "unit_label": rate_unit_label(rate_by) if rate_by else "p/unit",
        "currencies": currencies,
        "minimums": minimums,
        "rates": rates,
        "has_minimum": cost_has_minimum(minimums),
    }


def discover_value_detail_pairs(df: pd.DataFrame) -> list[tuple[str, str, str]]:
    pairs: list[tuple[str, str, str]] = []
    for col in df.columns:
        if not col.endswith(" Value"):
            continue
        cost_name = col[: -len(" Value")]
        details_col = f"{cost_name} Details"
        if details_col in df.columns:
            pairs.append((cost_name, col, details_col))
    return pairs


def format_validity_bracket(df: pd.DataFrame) -> str:
    """
    Build '(dd.mm.yyyy - dd.mm.yyyy)' from Line Item Eff/Exp dates
    when they are present (used for BAF Fee header).
    """
    if "Line Item Eff Date" not in df.columns or "Line Item Exp Date" not in df.columns:
        return ""

    pairs = (
        df[["Line Item Eff Date", "Line Item Exp Date"]]
        .dropna(how="all")
        .drop_duplicates()
    )
    if pairs.empty:
        return ""

    # Prefer a single shared window; otherwise use the first one
    eff = _format_date(pairs.iloc[0]["Line Item Eff Date"])
    exp = _format_date(pairs.iloc[0]["Line Item Exp Date"])
    if not eff and not exp:
        return ""
    if eff and exp:
        return f"({eff} - {exp})"
    return f"({eff or exp})"


def display_cost_name(source_name: str, df: pd.DataFrame | None = None) -> str:
    """Apply renames; BAF Fee also gets validity in brackets."""
    name = COST_DISPLAY_NAMES.get(source_name, source_name)
    if source_name == "Bunker Adjustment Factor" and df is not None:
        validity = format_validity_bracket(df)
        if validity:
            return f"{name} {validity}"
    return name


def build_value_detail_cost(
    df: pd.DataFrame,
    cost_name: str,
    value_col: str,
    details_col: str,
) -> dict:
    """Value/Details cost as a 2/3-col block (no Minimum unless present later)."""
    currencies: list = []
    minimums: list = []
    rates: list = []
    bases: list[str] = []

    for _, row in df.iterrows():
        currency, basis = parse_details(row.get(details_col))
        currencies.append(currency)
        minimums.append(None)
        rates.append(visible_number(row.get(value_col)))
        bases.append(basis or "")

    rate_by = _shared_or_first(bases)
    return {
        "cost_name": display_cost_name(cost_name, df),
        "rate_by": rate_by,
        "unit_label": rate_unit_label(rate_by) if rate_by else "p/unit",
        "currencies": currencies,
        "minimums": minimums,
        "rates": rates,
        "has_minimum": cost_has_minimum(minimums),
    }


def build_all_cost_columns(df: pd.DataFrame) -> list[dict]:
    """Base Rate first, then Value/Details rates. Drop fully empty costs."""
    costs: list[dict] = []
    if "Base Rate" in df.columns:
        costs.append(build_base_rate_cost(df))
    for cost_name, value_col, details_col in discover_value_detail_pairs(df):
        costs.append(build_value_detail_cost(df, cost_name, value_col, details_col))
    return [c for c in costs if not cost_is_empty(c)]


def cost_is_empty(cost: dict) -> bool:
    """True when the cost has no rates and no minimums on any lane."""
    has_rate = any(v is not None for v in cost.get("rates", []))
    has_min = any(v is not None for v in cost.get("minimums", []))
    return not has_rate and not has_min


def load_baf_rates(path: Path) -> pd.DataFrame:
    frames = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    if "Sheet1" in frames:
        df = frames["Sheet1"]
    elif "Mainleg rates" in frames:
        df = frames["Mainleg rates"]
    else:
        df = next(iter(frames.values()))
    return _strip_columns(df)


def load_baf_files(
    baf_paths: Path | str | list[Path | str] | None,
) -> tuple[pd.DataFrame | None, dict[str, str]]:
    """
    Load one or more BAF files into a combined dataframe.
    Returns (combined_df, uid -> source file name for rate card).
    """
    if baf_paths is None:
        return None, {}
    if isinstance(baf_paths, (str, Path)):
        paths = [Path(baf_paths)]
    else:
        paths = [Path(p) for p in baf_paths]

    frames: list[pd.DataFrame] = []
    uid_source: dict[str, str] = {}
    for path in paths:
        if not path.exists():
            continue
        df = load_baf_rates(path)
        df = df.copy()
        df["_baf_source"] = path.name
        if path.name.endswith("_processed.xlsx"):
            rate_card = path.name[: -len("_processed.xlsx")] + ".xlsx"
        else:
            rate_card = path.name
        for _, row in df.iterrows():
            uid = format_unique_identifier(row.get("Unique Identifier"))
            if uid and uid not in uid_source:
                uid_source[uid] = rate_card
        frames.append(df)

    if not frames:
        return None, {}
    return pd.concat(frames, ignore_index=True), uid_source


def collect_baf_uids(baf_df: pd.DataFrame) -> set[str]:
    uids: set[str] = set()
    for value in baf_df.get("Unique Identifier", []):
        uid = format_unique_identifier(value)
        if uid:
            uids.add(uid)
    return uids


def collect_main_uids(df: pd.DataFrame) -> set[str]:
    uids: set[str] = set()
    for value in df.get("Unique Identifier", []):
        uid = format_unique_identifier(value)
        if uid:
            uids.add(uid)
    return uids


def lane_involves_gb(row: pd.Series) -> bool:
    """True when Origin Country or Destination Country is GB."""
    origin = row.get("Origin Country")
    dest = row.get("Destination Country")
    origin_code = str(origin).strip().upper() if pd.notna(origin) else ""
    dest_code = str(dest).strip().upper() if pd.notna(dest) else ""
    return origin_code == "GB" or dest_code == "GB"


def extract_missing_baf_lanes(
    baf_df: pd.DataFrame,
    missing_uids: set[str],
    uid_source: dict[str, str],
    *,
    gb_only: bool | None = True,
) -> tuple[pd.DataFrame, list[str]]:
    """
    One representative BAF-file row per missing UID (earliest Eff Date).

    gb_only:
      True  -> only Origin/Destination country GB
      False -> only non-GB
      None  -> all missing UIDs
    """
    if not missing_uids or baf_df is None or baf_df.empty:
        return pd.DataFrame(), []

    work = baf_df.copy()
    work["_uid"] = work["Unique Identifier"].map(format_unique_identifier)
    work = work[work["_uid"].isin(missing_uids)]
    if work.empty:
        return pd.DataFrame(), []

    if "Line Item Eff Date" in work.columns:
        work = work.sort_values("Line Item Eff Date", kind="mergesort")
    work = work.drop_duplicates(subset=["_uid"], keep="first")

    if gb_only is True:
        work = work[work.apply(lane_involves_gb, axis=1)]
    elif gb_only is False:
        work = work[~work.apply(lane_involves_gb, axis=1)]

    if work.empty:
        return pd.DataFrame(), []

    rate_cards = [
        uid_source.get(uid, "BAF") for uid in work["_uid"].tolist()
    ]
    extra = work.drop(columns=["_uid", "_baf_source"], errors="ignore")
    return extra.reset_index(drop=True), rate_cards


def merge_main_with_missing_baf(
    main_df: pd.DataFrame,
    main_rate_card: str,
    extra_df: pd.DataFrame,
    extra_rate_cards: list[str],
) -> tuple[pd.DataFrame, list[str], list[bool]]:
    """
    Concatenate main lanes + missing BAF lanes.
    Returns (combined_df, rate_card_per_row, is_missing_baf_row flags).
    """
    main_cards = [main_rate_card] * len(main_df)
    main_flags = [False] * len(main_df)

    if extra_df is None or extra_df.empty:
        return main_df.copy(), main_cards, main_flags

    # Align columns: union of both, missing filled with NA
    combined = pd.concat([main_df, extra_df], ignore_index=True, sort=False)
    rate_cards = main_cards + list(extra_rate_cards)
    flags = main_flags + [True] * len(extra_df)
    return combined, rate_cards, flags


def build_shipment_rows_with_cards(
    df: pd.DataFrame,
    rate_cards: list[str],
) -> list[list]:
    """Like build_shipment_rows but with a per-row Rate Card."""
    rows = build_shipment_rows(df, rate_cards[0] if rate_cards else "")
    for i, row in enumerate(rows):
        if i < len(rate_cards):
            row[1] = rate_cards[i]
    return rows


def compute_baf_period_windows(baf_df: pd.DataFrame) -> list[dict]:
    """
    Build BAF validity windows from Line Item Eff/Exp dates.

    When multiple Eff dates exist, each period ends the day before the next
    Eff date (e.g. 2026-07-15 / 2026-12-31 with next Eff 2026-08-15
    -> BAF Fee (15.07.2026 - 14.08.2026)). The last period uses Exp date.
    """
    if "Line Item Eff Date" not in baf_df.columns:
        return []

    pairs = (
        baf_df[["Line Item Eff Date", "Line Item Exp Date"]]
        .dropna(subset=["Line Item Eff Date"])
        .drop_duplicates()
        .sort_values("Line Item Eff Date")
        .reset_index(drop=True)
    )
    if pairs.empty:
        return []

    windows: list[dict] = []
    for i, row in pairs.iterrows():
        eff = pd.Timestamp(row["Line Item Eff Date"])
        exp = (
            pd.Timestamp(row["Line Item Exp Date"])
            if pd.notna(row.get("Line Item Exp Date"))
            else None
        )
        if i + 1 < len(pairs):
            end = pd.Timestamp(pairs.iloc[i + 1]["Line Item Eff Date"]) - pd.Timedelta(
                days=1
            )
        else:
            end = exp

        eff_fmt = eff.strftime("%d.%m.%Y")
        end_fmt = end.strftime("%d.%m.%Y") if end is not None else ""
        if end_fmt:
            label = f"BAF Fee ({eff_fmt} - {end_fmt})"
        else:
            label = f"BAF Fee ({eff_fmt})"

        windows.append(
            {
                "eff": eff,
                "end": end,
                "label": label,
                "eff_key": row["Line Item Eff Date"],
            }
        )
    return windows


def build_baf_period_cost(
    baf_df: pd.DataFrame,
    window: dict,
    main_uids: list[str | None],
) -> dict:
    """One BAF Fee cost column for a validity window, mapped by Unique Identifier."""
    eff_key = window["eff_key"]
    period = baf_df[
        pd.to_datetime(baf_df["Line Item Eff Date"], errors="coerce")
        == pd.Timestamp(eff_key)
    ]

    by_uid: dict[str, dict] = {}
    for _, row in period.iterrows():
        uid = format_unique_identifier(row.get("Unique Identifier"))
        if not uid:
            continue
        currency, basis = parse_details(row.get("Bunker Adjustment Factor Details"))
        by_uid[uid] = {
            "currency": currency,
            "basis": basis or "",
            "rate": visible_number(row.get("Bunker Adjustment Factor Value")),
        }

    currencies: list = []
    rates: list = []
    bases: list[str] = []
    for uid in main_uids:
        hit = by_uid.get(uid) if uid else None
        if hit:
            currencies.append(hit["currency"])
            rates.append(hit["rate"])
            bases.append(hit["basis"])
        else:
            currencies.append(None)
            rates.append(None)
            bases.append("")

    rate_by = _shared_or_first(bases)
    return {
        "cost_name": window["label"],
        "rate_by": rate_by,
        "unit_label": rate_unit_label(rate_by) if rate_by else "p/unit",
        "currencies": currencies,
        "minimums": [None] * len(main_uids),
        "rates": rates,
        "has_minimum": False,
    }


def prolong_baf_costs_from_previous(costs: list[dict]) -> list[dict]:
    """
    For each later BAF period, if a lane rate is empty and the previous
    period has a value, copy currency + rate from the previous period.
    """
    if len(costs) < 2:
        return costs

    for idx in range(1, len(costs)):
        prev = costs[idx - 1]
        curr = costs[idx]
        for i, rate in enumerate(curr["rates"]):
            if rate is not None:
                continue
            prev_rate = prev["rates"][i]
            if prev_rate is None:
                continue
            curr["rates"][i] = prev_rate
            curr["currencies"][i] = prev["currencies"][i]
            # keep minimums aligned (BAF periods have none)
            curr["minimums"][i] = prev["minimums"][i]

        # Refresh shared rate_by / unit if still empty but we filled values
        if not curr["rate_by"] and prev["rate_by"]:
            curr["rate_by"] = prev["rate_by"]
            curr["unit_label"] = prev["unit_label"]

    return costs


def fill_first_baf_period_from_main(
    costs: list[dict],
    shipment_rows: list[list],
    main_df: pd.DataFrame,
) -> list[dict]:
    """
    Seed the earliest BAF validity column with main-rates BAF values
    where that column is empty (do not overwrite existing BAF-file values).
    """
    if not costs:
        return costs
    if "Bunker Adjustment Factor Value" not in main_df.columns:
        return costs

    main_by_uid: dict[str, dict] = {}
    for _, row in main_df.iterrows():
        uid = format_unique_identifier(row.get("Unique Identifier"))
        if not uid:
            continue
        currency, basis = parse_details(row.get("Bunker Adjustment Factor Details"))
        rate = visible_number(row.get("Bunker Adjustment Factor Value"))
        if rate is None:
            continue
        main_by_uid[uid] = {
            "currency": currency,
            "basis": basis or "",
            "rate": rate,
        }

    first = costs[0]
    filled = 0
    for i, ship_row in enumerate(shipment_rows):
        if first["rates"][i] is not None:
            continue
        uid = ship_row[2]
        hit = main_by_uid.get(uid) if uid else None
        if not hit:
            continue
        first["rates"][i] = hit["rate"]
        first["currencies"][i] = hit["currency"]
        filled += 1

    if filled and not first["rate_by"]:
        # Prefer shared basis from main BAF details
        bases = [v["basis"] for v in main_by_uid.values() if v.get("basis")]
        rate_by = _shared_or_first(bases)
        if rate_by:
            first["rate_by"] = rate_by
            first["unit_label"] = rate_unit_label(rate_by)

    if filled:
        print(
            f"  Filled first BAF period ({first['cost_name']}) "
            f"from main rates where empty: {filled} lanes"
        )
    return costs


def build_baf_sheet_costs(
    shipment_rows: list[list],
    baf_df: pd.DataFrame | None,
    main_df: pd.DataFrame | None = None,
) -> list[dict]:
    """
    BAF tab costs from all BAF files (one column per validity window,
    ordered by Eff date). Mapped by Unique Identifier.

    Then:
      - seed first period from main-rates BAF where empty (no overwrite)
      - prolong later periods from previous where empty
    """
    costs: list[dict] = []

    if baf_df is None or baf_df.empty:
        return costs

    lane_uids = [row[2] for row in shipment_rows]  # UNIQUE_IDENTIFIER
    for window in compute_baf_period_windows(baf_df):
        period_cost = build_baf_period_cost(baf_df, window, lane_uids)
        costs.append(period_cost)

    if main_df is not None:
        costs = fill_first_baf_period_from_main(costs, shipment_rows, main_df)

    costs = prolong_baf_costs_from_previous(costs)
    return [c for c in costs if not cost_is_empty(c)]


def write_matrix_sheet(
    ws,
    shipment_rows: list[list],
    cost_columns: list[dict],
    green_row_flags: list[bool] | None = None,
) -> None:
    """Write one matrix sheet (shipment details + cost blocks)."""
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    n_ship = len(SHIPMENT_HEADERS)
    flags = green_row_flags or [False] * len(shipment_rows)

    def style_header(cell):
        cell.font = header_font
        cell.alignment = center
        cell.fill = header_fill
        cell.border = border

    for col_idx, name in enumerate(SHIPMENT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        style_header(cell)
        ws.merge_cells(
            start_row=1, start_column=col_idx, end_row=3, end_column=col_idx
        )
        style_header(ws.cell(row=1, column=col_idx))

    next_col = n_ship + 1
    for cost in cost_columns:
        include_min = bool(cost.get("has_minimum"))
        block_width = 3 if include_min else 2
        start_col = next_col
        end_col = start_col + block_width - 1
        next_col = end_col + 1

        rate_by = cost["rate_by"]
        rate_by_label = f"Rate by: {rate_by}" if rate_by else "Rate by:"
        unit_label = cost["unit_label"] or "p/unit"

        cell = ws.cell(row=1, column=start_col, value=cost["cost_name"])
        style_header(cell)
        ws.merge_cells(
            start_row=1, start_column=start_col, end_row=1, end_column=end_col
        )
        style_header(ws.cell(row=1, column=start_col))

        cell = ws.cell(row=2, column=start_col, value=rate_by_label)
        style_header(cell)
        ws.merge_cells(
            start_row=2, start_column=start_col, end_row=2, end_column=end_col
        )
        style_header(ws.cell(row=2, column=start_col))

        sub_headers = [("Currency", start_col)]
        rate_col = start_col + 1
        if include_min:
            sub_headers.append(("MIN\nFlat", start_col + 1))
            rate_col = start_col + 2
        sub_headers.append((unit_label, rate_col))

        for label, col in sub_headers:
            cell = ws.cell(row=3, column=col, value=label)
            style_header(cell)

        for i, _ in enumerate(shipment_rows):
            excel_row = DATA_START_ROW + i
            currency = cost["currencies"][i]
            minimum = cost["minimums"][i]
            rate = cost["rates"][i]

            has_amount = rate is not None or minimum is not None
            cells = [
                ws.cell(
                    row=excel_row,
                    column=start_col,
                    value=_excel_value(currency) if has_amount else None,
                )
            ]
            if include_min:
                cells.append(
                    ws.cell(
                        row=excel_row,
                        column=start_col + 1,
                        value=_excel_value(minimum),
                    )
                )
            cells.append(
                ws.cell(
                    row=excel_row,
                    column=rate_col,
                    value=_excel_value(rate),
                )
            )
            if i < len(flags) and flags[i]:
                for cell in cells:
                    cell.fill = green_fill

    max_col = max(n_ship, next_col - 1)

    for i, ship_row in enumerate(shipment_rows):
        excel_row = DATA_START_ROW + i
        for col_idx, value in enumerate(ship_row, start=1):
            cell = ws.cell(row=excel_row, column=col_idx, value=_excel_value(value))
            if i < len(flags) and flags[i]:
                cell.fill = green_fill

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 30
    _autosize(ws, max_col)
    ws.freeze_panes = "A4"


ACCESSORIAL_LANES_PER_PART = 250

ACCESSORIAL_FIXED_HEADERS = [
    "lane #",
    "Currency",
    "MIN",
    "p/unit",
    "Basis",
]

# 5 Applies-if columns (values filled per rule; headers left blank under merged title)
ACCESSORIAL_APPLIES_WIDTH = 5

ACCESSORIAL_TAIL_HEADERS = [
    "Valid from",
    "Valid to",
]


def _strip_service(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def _city_display(value, strip_comma: bool) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if not text:
        return ""
    if strip_comma:
        return text.split(",", 1)[0].strip()
    return text


def _accessorial_cost_signature(row: pd.Series) -> tuple:
    return (
        visible_number(row.get("Base Rate")),
        None
        if pd.isna(row.get("Base Rate Currency"))
        else str(row.get("Base Rate Currency")).strip(),
        parse_minimum_amount(row.get("Minimum")),
        None
        if pd.isna(row.get("Base Rate Basis"))
        else str(row.get("Base Rate Basis")).strip(),
        _format_date(row.get("Line Item Eff Date")),
        _format_date(row.get("Line Item Exp Date")),
    )


def _basis_value(row: pd.Series) -> str:
    if pd.isna(row.get("Base Rate Basis")):
        return ""
    return str(row.get("Base Rate Basis")).strip()


def _lane_geo_key(row: pd.Series, *, strip_comma: bool = False) -> tuple:
    """Lane identity: From/To city + services. Keep state after comma for matching."""
    return (
        _city_display(row.get("From City"), strip_comma=strip_comma).casefold(),
        _strip_service(row.get("From Service")),
        _city_display(row.get("To City"), strip_comma=strip_comma).casefold(),
        _strip_service(row.get("To Service")),
    )


def dedupe_accessorial_lanes(
    df: pd.DataFrame,
) -> tuple[pd.DataFrame, list[bool], list[bool], list[bool]]:
    """
    Group by From City / From Service / To City / To Service (full city incl. state).

    - Identical costs -> keep one lane; strip From City after comma
    - Differing costs -> keep all; keep full From City; mark red
    - To City always strips after comma (even for red / collision lanes)
    - If stripping From City would collide distinct origins, keep full From City
    - For red groups: show Base Rate Basis only when it differs within the group

    Returns (work_df, strip_from_flags, red_flags, show_basis_flags).
    """
    work = df.copy().reset_index(drop=True)
    strip_from_flags = [True] * len(work)
    red_flags = [False] * len(work)
    show_basis_flags = [False] * len(work)
    keep_indexes: list[int] = []

    groups: dict[tuple, list[int]] = {}
    for i, row in work.iterrows():
        groups.setdefault(_lane_geo_key(row, strip_comma=False), []).append(i)

    for indexes in groups.values():
        if len(indexes) == 1:
            keep_indexes.append(indexes[0])
            continue

        sigs = [_accessorial_cost_signature(work.loc[i]) for i in indexes]
        unique_sigs = set(sigs)
        if len(unique_sigs) == 1:
            # full duplicates (incl. costs) -> keep first only, strip From City
            keep_indexes.append(indexes[0])
            strip_from_flags[indexes[0]] = True
            red_flags[indexes[0]] = False
        else:
            # same geo key, different costs -> keep all, keep From City, highlight red
            bases = {_basis_value(work.loc[i]) for i in indexes}
            basis_differs = len(bases) > 1
            for i in indexes:
                keep_indexes.append(i)
                strip_from_flags[i] = False
                red_flags[i] = True
                show_basis_flags[i] = basis_differs

    keep_indexes = sorted(keep_indexes)
    out = work.loc[keep_indexes].reset_index(drop=True)
    out_strip_from = [strip_from_flags[i] for i in keep_indexes]
    out_red = [red_flags[i] for i in keep_indexes]
    out_show_basis = [show_basis_flags[i] for i in keep_indexes]

    # Keep From City state when stripping would collide (Athens, AL vs Athens, GA)
    stripped_groups: dict[tuple, list[int]] = {}
    for idx in range(len(out)):
        if not out_strip_from[idx] or out_red[idx]:
            continue
        key = _lane_geo_key(out.loc[idx], strip_comma=True)
        stripped_groups.setdefault(key, []).append(idx)
    for idxs in stripped_groups.values():
        if len(idxs) > 1:
            for i in idxs:
                out_strip_from[i] = False

    return out, out_strip_from, out_red, out_show_basis


def build_accessorial_fee_name(kind: str, df: pd.DataFrame) -> str:
    """
    Pre-carriage Fee( D => CFS-CY, CFS-P)
    On-carriage Fee( CFS-CY, CFS-P => D)
    """
    if df is None or df.empty:
        return "Pre-carriage Fee" if kind == "pre" else "On-carriage Fee"

    preferred_order = ["CFS-CY", "CFS-P", "D"]

    def ordered_unique(values) -> list[str]:
        found = {_strip_service(v) for v in values if _strip_service(v)}
        ordered = [s for s in preferred_order if s in found]
        ordered.extend(sorted(found - set(preferred_order)))
        return ordered

    from_vals = ordered_unique(df["From Service"])
    to_vals = ordered_unique(df["To Service"])

    if kind == "pre":
        left = from_vals[0] if len(from_vals) == 1 else ", ".join(from_vals)
        right = ", ".join(to_vals)
        return f"Pre-carriage Fee( {left} => {right})"

    left = ", ".join(from_vals)
    right = to_vals[0] if len(to_vals) == 1 else ", ".join(to_vals)
    return f"On-carriage Fee( {left} => {right})"


def _applies_if_rules(kind: str, from_city: str, to_city: str) -> list[list]:
    """
    Two Applies-if rules (5 values each) for one lane.
    pre: SHIP_CITY / TO_SHIP_CITY
    on:  CUST_CITY / TO_CUST_CITY
    """
    if kind == "pre":
        return [
            ["Order tag", "SHIP_CITY", "starts with", from_city, "all items"],
            ["Order tag", "TO_SHIP_CITY", "starts with", to_city, "all items"],
        ]
    return [
        ["Order tag", "CUST_CITY", "starts with", from_city, "all items"],
        ["Order tag", "TO_CUST_CITY", "starts with", to_city, "all items"],
    ]


def _resolve_accessorial_city(value, *, strip_comma: bool) -> tuple[str, list[str]]:
    """
    Resolve city for display:
      1) spelling rename on full raw (keeps full target, e.g. "a, b, c")
      2) else strip after comma if requested, then spelling rename
      3) else stripped/raw city
    Spelling is the full target string — not common-rating parentheses.
    """
    raw = "" if pd.isna(value) else str(value).strip()
    if not raw:
        return "", []

    spelling = lookup_city_spelling(raw)
    if spelling is not None:
        return spelling, []

    base = _city_display(raw, strip_comma=strip_comma)
    spelling = lookup_city_spelling(base)
    if spelling is not None:
        return spelling, []

    return base, []


def build_accessorial_block(kind: str, df: pd.DataFrame) -> list[dict]:
    """
    Build Accessorial cost block(s) from pre-carriage or on-carriage sheet.
    Splits into parts of ACCESSORIAL_LANES_PER_PART lanes; lane # restarts each part.
    """
    if df is None or df.empty:
        return []

    work = df.copy()
    work.columns = [" ".join(str(c).strip().split()) for c in work.columns]

    # Apply city renames before dedupe so e.g. Machelen (Zulte) merges with Machelen
    if "From City" in work.columns:
        work["From City"] = [
            apply_city_rename("" if pd.isna(v) else str(v).strip())[0]
            for v in work["From City"]
        ]
    if "To City" in work.columns:
        work["To City"] = [
            apply_city_rename("" if pd.isna(v) else str(v).strip())[0]
            for v in work["To City"]
        ]

    work, strip_from_flags, red_flags, show_basis_flags = dedupe_accessorial_lanes(work)

    bases = [
        str(v).strip()
        for v in work.get("Base Rate Basis", pd.Series(dtype=object))
        if pd.notna(v) and str(v).strip()
    ]
    rate_by = _shared_or_first(bases)
    fee_name = build_accessorial_fee_name(kind, work)

    # Resolve display cities (strip rules + rename aliases) then common-rate format
    resolved: list[dict] = []
    for pos, (_, row) in enumerate(work.iterrows()):
        strip_from = strip_from_flags[pos]
        from_city, from_aliases = _resolve_accessorial_city(
            row.get("From City"), strip_comma=strip_from
        )
        to_city, to_aliases = _resolve_accessorial_city(
            row.get("To City"), strip_comma=True
        )
        resolved.append(
            {
                "from_city": from_city,
                "from_aliases": from_aliases,
                "to_city": to_city,
                "to_aliases": to_aliases,
                "from_service": _strip_service(row.get("From Service")),
                "to_service": _strip_service(row.get("To Service")),
                "currency": (
                    None
                    if pd.isna(row.get("Base Rate Currency"))
                    else str(row.get("Base Rate Currency")).strip() or None
                ),
                "min_amt": parse_minimum_amount(row.get("Minimum")),
                "rate": visible_number(row.get("Base Rate")),
                "basis": _basis_value(row) if show_basis_flags[pos] else None,
                "valid_from": _format_date(row.get("Line Item Eff Date")),
                "valid_to": _format_date(row.get("Line Item Exp Date")),
                "red": red_flags[pos],
            }
        )

    lane_index = build_lane_city_index(
        [
            (r["from_city"], r["from_service"], r["to_city"], r["to_service"])
            for r in resolved
        ]
    )

    lanes: list[dict] = []
    for r in resolved:
        from_key = city_match_key(r["from_city"])
        to_key = city_match_key(r["to_city"])
        from_city = format_city_with_aliases(
            r["from_city"],
            side="from",
            from_key=from_key,
            from_service=r["from_service"],
            to_key=to_key,
            to_service=r["to_service"],
            lane_index=lane_index,
            rename_aliases=r["from_aliases"],
        )
        to_city = format_city_with_aliases(
            r["to_city"],
            side="to",
            from_key=from_key,
            from_service=r["from_service"],
            to_key=to_key,
            to_service=r["to_service"],
            lane_index=lane_index,
            rename_aliases=r["to_aliases"],
        )
        lanes.append(
            {
                "currency": r["currency"],
                "min_amt": r["min_amt"],
                "rate": r["rate"],
                "basis": r["basis"],
                "applies": _applies_if_rules(kind, from_city, to_city),
                "valid_from": r["valid_from"],
                "valid_to": r["valid_to"],
                "red": r["red"],
            }
        )

    blocks: list[dict] = []
    part_size = ACCESSORIAL_LANES_PER_PART
    for part_i, start in enumerate(range(0, len(lanes), part_size), start=1):
        chunk = lanes[start : start + part_size]
        rows: list[list] = []
        row_red: list[bool] = []
        for lane_no, lane in enumerate(chunk, start=1):
            for applies in lane["applies"]:
                rows.append(
                    [
                        lane_no,
                        lane["currency"],
                        lane["min_amt"],
                        lane["rate"],
                        lane["basis"],
                        *applies,
                        lane["valid_from"],
                        lane["valid_to"],
                    ]
                )
                row_red.append(lane["red"])
        blocks.append(
            {
                "cost_name": f"{fee_name} /part {part_i}",
                "rate_by": rate_by,
                "rows": rows,
                "red_flags": row_red,
            }
        )
    return blocks


def build_accessorial_blocks(main_rates_path: Path) -> list[dict]:
    """Load pre-carriage / on-carriage sheets and build Accessorial cost blocks."""
    frames = pd.read_excel(main_rates_path, sheet_name=None, engine="openpyxl")
    by_lower = {str(k).strip().lower(): v for k, v in frames.items()}

    blocks: list[dict] = []
    pre = by_lower.get("pre-carriage")
    on_ = by_lower.get("on-carriage")

    if pre is not None:
        parts = build_accessorial_block("pre", pre)
        blocks.extend(parts)
        if parts:
            total_lanes = sum(
                (p["rows"][-1][0] if p["rows"] else 0) for p in parts
            )
            red_n = sum(sum(1 for f in p["red_flags"] if f) for p in parts)
            print(
                f"  Accessorial pre-carriage: {total_lanes} lanes in {len(parts)} part(s) "
                f"({sum(len(p['rows']) for p in parts)} applies-if rows"
                + (f", {red_n} red" if red_n else "")
                + ")"
            )
    if on_ is not None:
        parts = build_accessorial_block("on", on_)
        blocks.extend(parts)
        if parts:
            total_lanes = sum(
                (p["rows"][-1][0] if p["rows"] else 0) for p in parts
            )
            red_n = sum(sum(1 for f in p["red_flags"] if f) for p in parts)
            print(
                f"  Accessorial on-carriage: {total_lanes} lanes in {len(parts)} part(s) "
                f"({sum(len(p['rows']) for p in parts)} applies-if rows"
                + (f", {red_n} red" if red_n else "")
                + ")"
            )
    return blocks


def write_accessorial_sheet(ws, blocks: list[dict]) -> None:
    """
    Accessorial costs layout per cost block:
      row1: Cost Name (merged)
      row2: Rate by: ... (merged)
      row3: lane # | Currency | MIN | p/unit | Basis | Applies if (x5) | Valid from | Valid to
      row4+: data (2 applies-if rows per lane)
    """
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(cell):
        cell.font = header_font
        cell.alignment = center
        cell.fill = header_fill
        cell.border = border

    block_width = (
        len(ACCESSORIAL_FIXED_HEADERS)
        + ACCESSORIAL_APPLIES_WIDTH
        + len(ACCESSORIAL_TAIL_HEADERS)
    )
    next_col = 1

    for block in blocks:
        start = next_col
        end = start + block_width - 1
        next_col = end + 2

        rate_by = block.get("rate_by") or ""
        rate_label = f"Rate by: {rate_by}" if rate_by else "Rate by:"

        cell = ws.cell(row=1, column=start, value=block["cost_name"])
        style_header(cell)
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        style_header(ws.cell(row=1, column=start))

        cell = ws.cell(row=2, column=start, value=rate_label)
        style_header(cell)
        ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)
        style_header(ws.cell(row=2, column=start))

        # Row 3 headers
        col = start
        for name in ACCESSORIAL_FIXED_HEADERS:
            style_header(ws.cell(row=3, column=col, value=name))
            col += 1

        applies_start = col
        style_header(ws.cell(row=3, column=applies_start, value="Applies if"))
        for i in range(1, ACCESSORIAL_APPLIES_WIDTH):
            style_header(ws.cell(row=3, column=applies_start + i, value=None))
        ws.merge_cells(
            start_row=3,
            start_column=applies_start,
            end_row=3,
            end_column=applies_start + ACCESSORIAL_APPLIES_WIDTH - 1,
        )
        style_header(ws.cell(row=3, column=applies_start))
        col = applies_start + ACCESSORIAL_APPLIES_WIDTH

        for name in ACCESSORIAL_TAIL_HEADERS:
            style_header(ws.cell(row=3, column=col, value=name))
            col += 1

        red_flags = block.get("red_flags") or [False] * len(block["rows"])
        for r_i, row_vals in enumerate(block["rows"]):
            excel_row = DATA_START_ROW + r_i
            for c_i, value in enumerate(row_vals):
                cell = ws.cell(
                    row=excel_row,
                    column=start + c_i,
                    value=_excel_value(value),
                )
                if r_i < len(red_flags) and red_flags[r_i]:
                    cell.fill = red_fill

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    max_col = max(1, next_col - 2)
    _autosize(ws, max_col)
    ws.freeze_panes = "A4"


def write_result_matrix(
    shipment_rows: list[list],
    cost_columns: list[dict],
    output_path: Path,
    baf_shipment_rows: list[list] | None = None,
    baf_cost_columns: list[dict] | None = None,
    green_row_flags: list[bool] | None = None,
    missing_shipment_rows: list[list] | None = None,
    missing_cost_columns: list[dict] | None = None,
    accessorial_blocks: list[dict] | None = None,
) -> Path:
    """Write Result (+ optional BAF + Missing lanes + Accessorial costs)."""
    _sync_dirs()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Result"
    write_matrix_sheet(
        ws,
        shipment_rows,
        cost_columns,
        green_row_flags=green_row_flags,
    )

    if baf_shipment_rows is not None and baf_cost_columns is not None:
        baf_ws = wb.create_sheet("BAF")
        write_matrix_sheet(
            baf_ws,
            baf_shipment_rows,
            baf_cost_columns,
            green_row_flags=green_row_flags,
        )

    if missing_shipment_rows and missing_cost_columns is not None:
        miss_ws = wb.create_sheet("Missing lanes")
        write_matrix_sheet(
            miss_ws,
            missing_shipment_rows,
            missing_cost_columns,
        )

    if accessorial_blocks:
        acc_ws = wb.create_sheet("Accessorial costs")
        write_accessorial_sheet(acc_ws, accessorial_blocks)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _excel_value(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def _autosize(ws, max_col: int) -> None:
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(ws.max_row, 60) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                # account for wrapped MIN\nFlat
                max_len = max(max_len, max(len(line) for line in str(val).split("\n")))
        ws.column_dimensions[letter].width = min(max(10, max_len + 2), 42)


def build_result_matrix(
    main_rates_path: Path | str,
    output_path: Path | str | None = None,
    rate_card: str | None = None,
    baf_path: Path | str | list[Path | str] | None = None,
) -> Path:
    """
    Build result matrix xlsx from main rates (+ optional BAF file(s)).

    Flow:
      1) Load main rates
      2) Load all BAF files; find UIDs present in BAF but missing from main
      3) GB missing lanes -> Result (green); other missing -> "Missing lanes" tab
      4) Build BAF tab with period costs ordered by validity + prolong
    """
    _sync_dirs()
    main_rates_path = Path(main_rates_path)
    main_df = load_main_rates(main_rates_path)
    card = rate_card or rate_card_name_from_path(main_rates_path)

    baf_df, uid_source = load_baf_files(baf_path)

    green_flags: list[bool] = [False] * len(main_df)
    rate_cards = [card] * len(main_df)
    combined_df = main_df

    missing_shipment_rows: list[list] | None = None
    missing_cost_columns: list[dict] | None = None

    if baf_df is not None and not baf_df.empty:
        main_uids = collect_main_uids(main_df)
        baf_uids = collect_baf_uids(baf_df)
        missing_uids = baf_uids - main_uids
        if missing_uids:
            gb_df, gb_cards = extract_missing_baf_lanes(
                baf_df, missing_uids, uid_source, gb_only=True
            )
            other_df, other_cards = extract_missing_baf_lanes(
                baf_df, missing_uids, uid_source, gb_only=False
            )

            if not gb_df.empty:
                print(
                    f"  Missing BAF UIDs with GB: {len(gb_df)} "
                    "added to Result (green)"
                )
                combined_df, rate_cards, green_flags = merge_main_with_missing_baf(
                    main_df, card, gb_df, gb_cards
                )
            else:
                print(
                    f"  Missing BAF UIDs: {len(missing_uids)} "
                    "(none with GB — nothing added to Result)"
                )

            if not other_df.empty:
                print(
                    f"  Missing BAF UIDs without GB: {len(other_df)} "
                    'written to "Missing lanes" tab'
                )
                missing_shipment_rows = build_shipment_rows_with_cards(
                    other_df, other_cards
                )
                missing_cost_columns = build_all_cost_columns(other_df)

    shipment_rows = build_shipment_rows_with_cards(combined_df, rate_cards)
    cost_columns = build_all_cost_columns(combined_df)

    baf_costs = build_baf_sheet_costs(shipment_rows, baf_df, main_df=main_df)
    accessorial_blocks = build_accessorial_blocks(main_rates_path)

    if output_path is None:
        output_path = result_output_path(main_rates_path, rate_card=card)
    else:
        output_path = Path(output_path)

    return write_result_matrix(
        shipment_rows,
        cost_columns,
        output_path,
        baf_shipment_rows=shipment_rows if baf_costs else None,
        baf_cost_columns=baf_costs if baf_costs else None,
        green_row_flags=green_flags,
        missing_shipment_rows=missing_shipment_rows,
        missing_cost_columns=missing_cost_columns,
        accessorial_blocks=accessorial_blocks or None,
    )


if __name__ == "__main__":
    default_main = PROCESSING_DIR / "K&N Ocean UK rates_processed.xlsx"
    if not default_main.exists():
        raise SystemExit(f"Main rates file not found: {default_main}")
    default_baf = PROCESSING_DIR / "BAF update august K&N LCL (2)_processed.xlsx"
    out = build_result_matrix(
        default_main,
        baf_path=default_baf if default_baf.exists() else None,
    )
    print(f"Wrote {out}")
